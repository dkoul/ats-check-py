"""File parsers for resume and job description files."""

from pathlib import Path
import PyPDF2
from docx import Document


def parse_resume(file_path: Path) -> str:
    """
    Extract text from resume file (PDF or DOCX).

    Args:
        file_path: Path to resume file

    Returns:
        Extracted text content
    """
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        return parse_pdf(file_path)
    elif suffix == ".docx":
        return parse_docx(file_path)
    else:
        raise ValueError(f"Unsupported file format: {suffix}")


def parse_pdf(file_path: Path) -> str:
    """
    Extract text from PDF file.

    Args:
        file_path: Path to PDF file

    Returns:
        Extracted text content
    """
    text = []

    with open(file_path, "rb") as file:
        reader = PyPDF2.PdfReader(file)

        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text.append(page_text)

    return "\n".join(text)


def parse_docx(file_path: Path) -> str:
    """
    Extract text from DOCX file.

    Args:
        file_path: Path to DOCX file

    Returns:
        Extracted text content
    """
    doc = Document(file_path)
    text = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text.append(cell.text)

    return "\n".join(text)


def parse_text_file(file_path: Path) -> str:
    """
    Read text from a plain text file.

    Args:
        file_path: Path to text file

    Returns:
        File content
    """
    with open(file_path, "r", encoding="utf-8") as file:
        return file.read()
